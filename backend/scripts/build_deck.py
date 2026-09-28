"""Build RootIQ_Pitch_Deck_Final.pptx from measured RESULTS + docs/deck screenshots."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
DECK = ROOT / "docs" / "deck"
OUT = ROOT / "docs" / "RootIQ_Pitch_Deck_Final.pptx"

# Measured — docs/RESULTS.md (Day 12 /api/runs summary)
AVG_TTRCA = "12.3 s"
TOP1 = "9/9"
NOISE = "96.3%"
APPROVAL = "100%"


def add_title(slide, text, top=0.4):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(9), Inches(0.7))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)


def add_body(slide, lines, top=1.3, size=16):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(top), Inches(9), Inches(5))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        p.space_after = Pt(8)


def main():
    DECK.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 Title
    s = prs.slides.add_slide(blank)
    add_title(s, "RootIQ — Find the cause before it becomes an outage")
    add_body(
        s,
        [
            "Alert storm → one incident → evidence-backed RCA → approve-only remediation",
            "Venture X · Infrastructure & Cloud",
            "Source of numbers: docs/RESULTS.md  ·  Architecture: docs/ARCHITECTURE.md",
        ],
        top=1.4,
        size=18,
    )

    # 2 Problem
    s = prs.slides.add_slide(blank)
    add_title(s, "The NOC problem")
    add_body(
        s,
        [
            "Dozens of alerts across router, DNS, and app — engineer correlating by hand",
            "Black-box tools hide why; auto-remediation without approval is a non-starter",
            "RootIQ: one correlated incident, ranked cause with evidence, human gate",
        ],
    )

    # 3 Promise
    s = prs.slides.add_slide(blank)
    add_title(s, "Promise")
    add_body(
        s,
        [
            "Under 60 seconds from inject to clear root cause",
            "Every remediation requires engineer approve/reject (audit logged)",
            "Works sim-first; live EVE-NG lab when venue network is ready",
        ],
    )

    # 4–7 placeholders for story arc (kept short)
    for title, bullets in [
        ("Lab topology", ["R1 · SW1/SW2 · APP-01 (DNS+Web) · COLLECTOR-01", "SNMP / ICMP / DNS / HTTP / Syslog → normalized events"]),
        ("How it works", ["Detect → Correlate → Rank (graph) → Explain (grounded) → Act (whitelist)"]),
        ("Modes", ["sim = in-process simulator (default / backup)", "live = EVE + collector + lab agent"]),
        ("Security", ["Read-oriented collect · fixed action whitelist · audit log · private deploy"]),
    ]:
        s = prs.slides.add_slide(blank)
        add_title(s, title)
        add_body(s, bullets)

    # 8 Demo — real screenshot
    s = prs.slides.add_slide(blank)
    add_title(s, "Demo — live map during incident")
    img = DECK / "08-demo-incident-map.png"
    if img.exists():
        s.shapes.add_picture(str(img), Inches(0.4), Inches(1.2), width=Inches(12.5))
    else:
        add_body(s, ["[screenshot missing — run frontend capture]"])

    # 9 Value — Measured in our lab
    s = prs.slides.add_slide(blank)
    add_title(s, "Measured in our lab (not pilot targets)")
    add_body(
        s,
        [
            f"Time to correlated RCA:  target < 60 s   →   measured {AVG_TTRCA} (avg of 9 runs)",
            f"Root ranking:             target Top 3    →   Top-1 correct {TOP1}",
            f"Alert noise reduction:    target 30%      →   measured {NOISE}",
            f"Human approval before remediation: 100%   →   {APPROVAL} (Audit)",
            "Source: GET /api/runs after backend/scripts/day12_measure.py → docs/RESULTS.md",
        ],
        size=18,
    )

    # Explainable
    s = prs.slides.add_slide(blank)
    add_title(s, "Why RootIQ is explainable")
    img2 = DECK / "explain-candidates.png"
    if img2.exists():
        s.shapes.add_picture(str(img2), Inches(0.4), Inches(1.15), width=Inches(12.5))
    add_body(
        s,
        [
            "CandidateRanking with score components · EvidenceList of measured metrics",
            "Grounded explanation: numbers must exist in evidence (LLM optional, never invents)",
        ],
        top=6.5,
        size=14,
    )

    # Architecture
    s = prs.slides.add_slide(blank)
    add_title(s, "Architecture")
    add_body(
        s,
        [
            "EVE lab → Collector → FastAPI ingest → State/Detector → Correlator → RCA+Graph → Explanation",
            "→ WebSocket → React dashboard → approve/reject → Action service → Lab agent → R1",
            "PostgreSQL when Docker is up  ·  Full Mermaid: docs/ARCHITECTURE.md",
        ],
        size=17,
    )
    healthy = DECK / "01-healthy-map.png"
    if healthy.exists():
        s.shapes.add_picture(str(healthy), Inches(7.2), Inches(3.2), width=Inches(5.5))

    # Close
    s = prs.slides.add_slide(blank)
    add_title(s, "RootIQ")
    box = s.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12), Inches(1))
    p = box.text_frame.paragraphs[0]
    p.text = "Find the cause before it becomes an outage."
    p.font.size = Pt(32)
    p.alignment = PP_ALIGN.CENTER
    p.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    add_body(s, [f"Measured: {AVG_TTRCA} RCA · {TOP1} top-1 · {NOISE} noise ↓ · {APPROVAL} approval"], top=4.0, size=18)

    prs.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
